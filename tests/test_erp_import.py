from __future__ import annotations

import json
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from auditpaper_agent.sensing.erp import diagnose_erp_export, import_erp_export, load_erp_package
from auditpaper_agent.suite import run_erp_workpaper_suite
from auditpaper_agent.workpaper_catalog import assert_workbook_clean_room


def _runtime_dir() -> Path:
    path = Path(__file__).resolve().parents[1] / "test_runtime" / uuid4().hex
    path.mkdir(parents=True, exist_ok=True)
    return path


def _fixture(name: str) -> Path:
    return Path(__file__).resolve().parent / "fixtures" / "erp_exports" / name


def test_sap_mapping_requires_confirmation_before_package_creation() -> None:
    export_dir = _fixture("sap")
    manifest = diagnose_erp_export(export_dir, provider="sap")

    assert not manifest.confirmed
    assert not manifest.blocking_issues
    assert {"trial_balance", "general_ledger", "customers", "suppliers", "inventory", "fixed_assets", "sales", "purchase"}.issubset(
        {table.source_type for table in manifest.tables}
    )
    tb_mapping = next(table for table in manifest.tables if table.source_type == "trial_balance")
    assert tb_mapping.header_row == 2
    assert tb_mapping.row_count == 5

    blocked = import_erp_export(export_dir, _runtime_dir() / "sap_case", provider="sap", confirm_mapping=False)

    assert not blocked.success
    assert blocked.manifest_path.exists()
    assert blocked.package_path is None
    assert "字段映射尚未确认" in blocked.errors[0]


def test_sap_import_builds_standardized_main_cycle_package() -> None:
    case_dir = _runtime_dir() / "sap_case"
    result = import_erp_export(_fixture("sap"), case_dir, provider="sap", confirm_mapping=True)

    assert result.success
    assert result.package_path and result.package_path.exists()
    package = load_erp_package(case_dir)

    assert package.mapping_confirmed
    assert package.provider == "sap"
    assert len(package.trial_balance) == 5
    assert len(package.general_ledger) == 5
    assert len(package.customers) == 2
    assert len(package.suppliers) == 2
    assert len(package.inventory) == 2
    assert len(package.fixed_assets) == 2
    assert package.bank_journal
    assert package.general_ledger[2].credit == 8_600_000
    assert package.trial_balance[4].ending_credit == 8_600_000


def test_yonyou_import_normalizes_direction_and_suite_generates() -> None:
    case_dir = _runtime_dir() / "yonyou_case"
    import_result = import_erp_export(_fixture("yonyou"), case_dir, provider="yonyou", confirm_mapping=True)
    assert import_result.success
    package = load_erp_package(case_dir)

    revenue = next(row for row in package.general_ledger if row.account_code == "600101")
    payable = next(row for row in package.trial_balance if row.account_code == "220201")
    assert revenue.debit == 0
    assert revenue.credit == 8_600_000
    assert payable.ending_balance == -900_000

    suite_dir = _runtime_dir() / "suite"
    suite = run_erp_workpaper_suite(case_dir, output_dir=suite_dir)

    assert suite.success
    assert set(suite.workbooks) == {"A10", "C", "D10", "E20", "EXP10", "F10", "K10", "N10", "U10"}
    assert suite.zip_path and suite.zip_path.exists()
    manifest = json.loads((suite_dir / "suite_manifest.json").read_text(encoding="utf-8"))
    assert manifest["input_mode"] == "erp_standardized_package"
    assert manifest["mapping_confirmed"] is True

    c_wb = load_workbook(suite.workbooks["C"], data_only=False)
    try:
        assert c_wb["Lead"]["C2"].value == "yonyou"
        assert "货币资金" in c_wb["Lead"]["B1"].value
    finally:
        c_wb.close()

    for workbook in suite.workbooks.values():
        assert_workbook_clean_room(workbook)


def test_unconfirmed_erp_package_is_blocked_from_suite_generation() -> None:
    case_dir = _runtime_dir() / "blocked_case"
    blocked = import_erp_export(_fixture("sap"), case_dir, provider="sap", confirm_mapping=False)
    assert not blocked.success

    suite = run_erp_workpaper_suite(case_dir)

    assert not suite.success
    assert suite.errors
